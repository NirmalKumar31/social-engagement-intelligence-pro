"""Export the Power BI star schema and write the PBIP project, Fabric notebook and Excel fallback."""
from __future__ import annotations

import argparse
import glob
import os
import json
import shutil
import uuid

import numpy as np
import pandas as pd

from . import config as C
from .narrative import facts
from .pbi_report import build_report

OUT = C.POWERBI
DATA = OUT / "data"
NAME = "Feed_Ranking_Launch_Review"
DISPLAY = "Feed Ranking Launch Review"
SEGMENTS = ["age_group", "region", "device", "account_type"]
SUMS = ["sessions", "minutes", "likes", "comments", "dms_sent", "shares", "posts_created", "negative_feedback",
        "ad_impressions", "ad_clicks", "ad_revenue"]


def export_data() -> dict[str, pd.DataFrame]:
    users = pd.read_parquet(C.RAW / "users.parquet")
    daily = pd.read_parquet(C.RAW / "daily_activity.parquet")
    ux = pd.read_parquet(C.PROCESSED / "user_experiment.parquet")
    tables: dict[str, pd.DataFrame] = {}

    days = np.arange(-C.PRE_DAYS, C.EXP_DAYS)
    dates = pd.Timestamp(C.START_DATE) + pd.to_timedelta(days, unit="D")
    tables["dim_date"] = pd.DataFrame({
        "date": dates, "experiment_day": days, "period": np.where(days < 0, "Pre-period", "Experiment"),
        "experiment_week": np.where(days < 0, 0, days // 7 + 1), "weekday": dates.day_name(),
        "is_weekend": (dates.dayofweek >= 5).astype(int)})
    tables["dim_arm"] = pd.DataFrame({"arm": C.ARMS, "arm_label": [C.ARM_LABELS[a] for a in C.ARMS],
                                      "is_treatment": [0, 1, 1], "sort_order": [1, 2, 3],
                                      "color_hex": [C.ARM_COLORS[a] for a in C.ARMS]})

    # Grain: date x arm x segment, with assigned-user counts as denominators
    d = daily.merge(users[["user_id", "arm"] + SEGMENTS], on="user_id")
    agg = d.groupby(["day", "arm"] + SEGMENTS)[SUMS].sum()
    agg["active_users"] = d.groupby(["day", "arm"] + SEGMENTS).size()
    assigned = users.groupby(["arm"] + SEGMENTS).size().rename("users_assigned").reset_index()
    grid = assigned.merge(pd.DataFrame({"day": days}), how="cross")
    fact = grid.merge(agg.reset_index(), on=["day", "arm"] + SEGMENTS, how="left").fillna(0)
    fact["date"] = pd.Timestamp(C.START_DATE) + pd.to_timedelta(fact.day, unit="D")
    fact["friend_interactions"] = fact.comments + fact.dms_sent
    fact = fact.drop(columns="day")
    int_cols = [c for c in SUMS + ["active_users", "users_assigned", "friend_interactions"] if c not in ("minutes", "ad_revenue")]
    fact[int_cols] = fact[int_cols].astype(int)
    fact["minutes"] = fact.minutes.round(1)
    fact["ad_revenue"] = fact.ad_revenue.round(4)
    tables["fact_daily"] = fact

    keep = ["user_id", "arm", "age_group", "region", "device", "account_type", "acquisition_channel", "tenure_days",
            "friend_count", "is_new_user", "churned", "pre_minutes_per_day", "exp_minutes_per_day",
            "pre_active_days_per_day", "exp_active_days_per_day", "exp_friend_interactions_per_day",
            "pre_ad_revenue_per_day", "exp_ad_revenue_per_day", "late_ad_revenue_per_day"]
    fu = ux[keep].copy()
    sc = pd.read_parquet(C.PROCESSED / "churn_scores.parquet")[["user_id", "churn_probability", "risk_tier", "risk_decile",
                                                                "actual_churn_next_14d"]]
    ca = pd.read_parquet(C.PROCESSED / "user_cate.parquet")[["user_id", "cate_ai_feed", "cate_ai_feed_social", "policy_arm"]]
    fu = fu.merge(sc, on="user_id", how="left").merge(ca, on="user_id", how="left")
    fu["risk_tier"] = fu.risk_tier.fillna("Not scored (inactive at day 13)")
    fu["policy_arm"] = fu.policy_arm.fillna("Not in held-out fold")
    tables["fact_user"] = fu.round(5)

    topic = pd.read_csv(C.TABLES / "topic_info.csv")
    tx = pd.read_csv(C.TABLES / "topic_exposure_by_arm.csv").merge(topic[["topic_id", "label"]], on="topic_id")
    simple = {
        "metric_results": "metric_results.csv", "segment_results": "segment_results.csv", "daily_lift": "daily_lift.csv",
        "sequential_monitoring": "sequential_monitoring.csv", "power_mde": "power_mde.csv", "balance_smd": "balance_smd.csv",
        "bayesian_results": "bayesian_results.csv", "km_survival": "km_survival.csv", "cox_hazard_ratios": "cox_hazard_ratios.csv",
        "churn_model_metrics": "churn_model_metrics.csv", "churn_shap_importance": "churn_shap_importance.csv",
        "retention_profit_curve": "retention_profit_curve.csv", "policy_values": "policy_values.csv",
        "qini_curves": "qini_curves.csv", "cate_deciles": "cate_deciles.csv", "topic_info": "topic_info.csv",
        "content_exposure_summary": "content_exposure_summary.csv", "sentiment_model_metrics": "sentiment_model_metrics.csv",
        "virality_playbook": "virality_playbook.csv", "virality_shap_importance": "virality_shap_importance.csv",
        "virality_weekly": "virality_weekly.csv", "impact_fan": "impact_fan.csv", "impact_distribution": "impact_distribution.csv",
        "impact_monte_carlo": "impact_monte_carlo.csv", "impact_sensitivity": "impact_sensitivity.csv",
    }
    for t, f in simple.items():
        tables[t] = pd.read_csv(C.TABLES / f)
    tables["topic_exposure_by_arm"] = tx
    tables["daily_lift"]["date"] = pd.to_datetime(tables["daily_lift"]["date"])
    tables["virality_weekly"]["week"] = pd.to_datetime(tables["virality_weekly"]["week"])
    for t in ["metric_results", "segment_results"]:
        for c in tables[t].columns:
            if pd.api.types.is_bool_dtype(tables[t][c]):
                tables[t][c] = tables[t][c].astype(int)
    tables["metric_results"]["metric_label"] = tables["metric_results"]["label"]
    add_display_columns(tables)

    if DATA.exists():
        shutil.rmtree(DATA)
    DATA.mkdir(parents=True)
    for t, df in tables.items():
        out = df.copy()
        for c in out.columns:
            if pd.api.types.is_datetime64_any_dtype(out[c]):
                out[c] = out[c].dt.strftime("%Y-%m-%d")
        out.to_csv(DATA / f"{t}.csv", index=False)
    return tables


# Readable columns for the report. Power BI column names are case-insensitive, so these must not
# collide with existing ones.
SHORT_METRIC = {"minutes": "Minutes / user", "active_days": "DAU rate", "sessions": "Sessions", "likes": "Likes",
                "shares": "Shares", "posts_created": "Posts created", "ad_revenue": "Ad revenue",
                "churned": "Churn (28-day)", "friend_interactions": "Friend interactions",
                "neg_feedback_rate": "Negative feedback", "ad_ctr": "Ad CTR"}
SHORT_POLICY = {"Ship nothing (control)": "Ship nothing", "Ship AI feed to all": "AI feed → all",
                "Ship AI + friend boost to all": "Friend boost → all", "Rule: AI feed <35, friend boost 35+": "Age rule (<35 / 35+)",
                "Model-targeted (DR-learner)": "Model-targeted"}
SHORT_DRIVER = {"Monthly base churn 2% / 6%": "Base churn 2–6%", "Eligible MAU 20M / 60M": "MAU 20–60M",
                "Revenue per user -30% / +30%": "Revenue/user ±30%",
                "AI-feed revenue lift at 95% CI low / high": "AI-feed revenue lift (CI)",
                "AI-feed churn HR at 95% CI low / high": "AI-feed churn HR (CI)",
                "Friend-boost churn HR at 95% CI low / high": "Friend-boost churn HR (CI)"}
FEATURE_NAMES = {"kw_avg_avg": "Keyword avg shares", "kw_max_avg": "Keyword max-avg shares", "kw_min_avg": "Keyword min-avg shares",
                 "kw_avg_max": "Keyword avg-max shares", "kw_max_max": "Keyword max shares", "kw_avg_min": "Keyword avg-min shares",
                 "self_reference_min_shares": "Linked-article min shares", "self_reference_avg_sharess": "Linked-article avg shares",
                 "self_reference_max_shares": "Linked-article max shares", "is_weekend": "Published on weekend",
                 "global_subjectivity": "Text subjectivity", "n_unique_tokens": "Unique-word ratio",
                 "n_non_stop_unique_tokens": "Unique non-stop words", "num_hrefs": "Number of links", "num_imgs": "Number of images",
                 "num_videos": "Number of videos", "n_tokens_content": "Article length", "n_tokens_title": "Title length",
                 "average_token_length": "Avg word length", "data_channel_is_entertainment": "Channel: entertainment",
                 "data_channel_is_world": "Channel: world", "data_channel_is_tech": "Channel: tech",
                 "data_channel_is_socmed": "Channel: social media", "global_rate_positive_words": "Positive-word rate",
                 "LDA_00": "Topic LDA 0", "LDA_01": "Topic LDA 1", "LDA_02": "Topic LDA 2", "LDA_03": "Topic LDA 3", "LDA_04": "Topic LDA 4",
                 "recency_days": "Days since last active", "minutes_w0": "Minutes, exp. week 1", "minutes_w1": "Minutes, exp. week 2",
                 "minutes_w-1": "Minutes, pre week 2", "minutes_w-2": "Minutes, pre week 1", "active_days_w0": "Active days, exp. week 1",
                 "active_days_w1": "Active days, exp. week 2", "active_days_w-1": "Active days, pre week 2",
                 "active_days_w-2": "Active days, pre week 1", "log_followers": "Followers (log)", "like_rate": "Likes per minute",
                 "minutes_per_session": "Minutes per session", "tenure_days": "Account tenure", "following_count": "Following count",
                 "friend_count": "Friend count", "ctr": "Ad CTR", "engagement_trend": "Engagement trend",
                 "activity_drop": "Activity drop", "friend_interactions_rate": "Friend interaction rate",
                 "neg_feedback_rate": "Negative feedback rate", "posts_total": "Posts created", "arm": "Experiment arm",
                 "age_group": "Age group", "region": "Region", "device": "Device", "account_type": "Account type",
                 "acquisition_channel": "Acquisition channel", "is_new_user": "New user"}
SEGMENT_DIM = {"age_group": "Age", "region": "Region", "device": "Device", "account_type": "Account type",
               "is_new_user": "New user", "acquisition_channel": "Channel"}


def add_display_columns(tables: dict[str, pd.DataFrame]):
    arm_label = tables["dim_arm"].set_index("arm").arm_label
    mr = tables["metric_results"]
    mr["Arm name"] = mr.arm.map(arm_label)
    mr["Metric name"] = mr.metric.map(SHORT_METRIC)
    mr["Lift"], mr["CI low"], mr["CI high"] = mr.rel_lift, mr.rel_ci_low, mr.rel_ci_high
    mr["p-value (BH)"] = mr.p_adj_bh
    mr["Variance removed"] = mr.variance_reduction
    mr["Result"] = mr.verdict.map({"win": "✔ Win", "harm": "✖ Harm", "neutral": "– Neutral"})
    sr = tables["segment_results"]
    sr["Arm name"] = sr.arm.map(arm_label)
    sr["Dimension name"] = sr.dimension.map(SEGMENT_DIM)
    sr["Segment value"] = np.where(sr.dimension == "is_new_user", sr.segment.map({"0": "Established", "1": "New (<90d)"}),
                                   sr.segment)
    sr["Minutes lift %"] = sr.minutes_lift
    sr["Churn change (pp)"] = sr.churn_diff_pp
    sr["Verdict"] = sr.status.map({"churn risk": "⚠ Churn risk", "clear win": "✔ Clear win", "inconclusive": "– Inconclusive"})
    pv = tables["policy_values"]
    pv["Policy name"] = pv.policy.map(SHORT_POLICY).fillna(pv.policy)
    sens = tables["impact_sensitivity"]
    sens["Driver"] = sens.assumption.map(SHORT_DRIVER).fillna(sens.assumption)
    sens["Low case ($M)"], sens["High case ($M)"] = sens.low_musd, sens.high_musd
    for t in ("virality_shap_importance", "churn_shap_importance"):
        tables[t]["Feature name"] = tables[t].feature.map(FEATURE_NAMES).fillna(tables[t].feature.str.replace("_", " "))
    sm = tables["sentiment_model_metrics"]
    sm["Model name"] = sm.model.str.replace(" (lexicon baseline)", "", regex=False).str.replace(" (ours)", "", regex=False)
    sm["Macro recall"], sm["Macro F1"], sm["Accuracy %"] = sm.macro_recall, sm.macro_f1, sm.accuracy
    mc = tables["impact_monte_carlo"]
    mc["Scenario"] = mc.comparison
    mc["P5 ($M)"], mc["Median ($M)"], mc["P95 ($M)"], mc["P(> 0)"] = mc.p05_musd, mc.p50_musd, mc.p95_musd, mc.prob_positive
    ce = tables["content_exposure_summary"]
    ce["Negative share"], ce["Topic diversity"] = ce.negative_share, ce.topic_entropy
    tables["topic_exposure_by_arm"]["Topic"] = tables["topic_exposure_by_arm"]["label"].str.title()
    tables["topic_exposure_by_arm"]["Impression share"] = tables["topic_exposure_by_arm"]["share"]
    pb = tables["virality_playbook"]
    pb["Playbook view"] = pb.dimension.map({"channel": "Channel", "weekday": "Weekday", "images_bucket": "Images",
                                            "videos_bucket": "Videos", "title_sentiment_bucket": "Title sentiment"})
    pb["Level name"] = pb.level.astype(str).str.title()
    pb["Viral rate"] = pb.viral_rate
    dec = json.loads((C.REPORTS / "experiment_summary.json").read_text())["decision"]
    icon = {"PASS": "✔ PASS", "FAIL": "✖ FAIL", "INCONCLUSIVE": "◐ INCONCLUSIVE"}
    rows = []
    for arm, a in dec["arms"].items():
        for metric, g in a["guardrails"].items():
            rows.append({"arm": arm, "Arm name": arm_label[arm], "Guardrail": SHORT_METRIC[metric] + (" ★" if g["severity"] == "critical" else ""),
                         "Observed change": g["rel_change"], "Harm bound": g["harm_bound"], "margin": g["margin"],
                         "Allowed harm": f"{'↑' if g['direction'] == 'increase' else '↓'} ≤ {g['margin']:.0%}",
                         "Status": icon[g["status"]], "Arm decision": a["state"]})
    tables["decision_guardrails"] = pd.DataFrame(rows)
    sc = pd.read_csv(C.TABLES / "impact_scenarios.csv")
    sc["Effect persistence"] = sc.persistence.str.capitalize()
    sc["Scenario"] = sc.comparison
    sc["Median ($M)"] = sc.p50_musd
    sc["90% range"] = [f"{'-' if lo < 0 else ''}${abs(lo):,.0f}M to {'-' if hi < 0 else ''}${abs(hi):,.0f}M"
                       for lo, hi in zip(sc.p05_musd, sc.p95_musd)]
    sc["Share of draws > 0"] = sc.prob_positive
    tables["impact_scenarios"] = sc
    mde = tables["power_mde"]
    mde["Method name"] = mde.method.map({"raw": "Raw t-test", "cuped": "With CUPED"})
    mde["Minimum detectable lift"] = mde.mde_rel


DISPLAY_FORMATS = {"Lift": "+0.0%;-0.0%;0.0%", "CI low": "+0.0%;-0.0%;0.0%", "CI high": "+0.0%;-0.0%;0.0%",
                   "p-value (BH)": "0.0000", "Variance removed": "0%", "Minutes lift %": "+0.0%;-0.0%;0.0%",
                   "Churn change (pp)": "+0.00;-0.00;0.00", "Low case ($M)": "$#,0", "High case ($M)": "$#,0",
                   "Macro recall": "0.000", "Macro F1": "0.000", "Accuracy %": "0.0%", "P5 ($M)": "$#,0",
                   "Median ($M)": "$#,0", "P95 ($M)": "$#,0", "P(> 0)": "0.00%", "Negative share": "0.0%",
                   "Topic diversity": "0.00", "Impression share": "0.0%", "Viral rate": "0%",
                   "Minimum detectable lift": "0.0%", "uplift_vs_control_usd": "$0.00", "p50": "$#,0", "bin_mid_musd": "$#,0",
                   "mean_abs_shap": "0.000", "profit_at_scale_musd": "$#,0.0", "predicted_cate": "$0.00",
                   "observed_uplift": "$0.00", "cumulative_gain": "#,0", "churn_diff_pp": "+0.00;-0.00;0.00",
                   "Observed change": "+0.0%;-0.0%;0.0%", "Harm bound": "+0.0%;-0.0%;0.0%", "Share of draws > 0": "0%"}

PCT_COLS = {"rel_lift", "rel_ci_low", "rel_ci_high", "minutes_lift", "minutes_ci_low", "minutes_ci_high", "friend_int_lift",
            "revenue_lift", "variance_reduction", "mde_rel", "survival", "ci_low_surv", "churn_probability", "share",
            "negative_share", "share_of_draws", "viral_rate", "recall", "precision", "p_improves", "fraction_targeted",
            "share_ai_feed", "share_ai_feed_social", "share_control", "prob_positive", "share_negative_pred",
            "share_positive_pred", "macro_recall", "macro_f1", "accuracy", "churn_control", "churn_treatment",
            "topic_entropy_rel_change", "negative_share_rel_change", "top_topic_share", "share_35_plus"}


def _dtype(s: pd.Series) -> tuple[str, str]:
    """Model and Power Query types for a column (works with pandas 3 string dtypes)."""
    t = pd.api.types
    if t.is_datetime64_any_dtype(s):
        return "dateTime", "type date"
    if t.is_bool_dtype(s):
        return "int64", "Int64.Type"
    if t.is_integer_dtype(s):
        return "int64", "Int64.Type"
    if t.is_float_dtype(s):
        return "double", "type number"
    return "string", "type text"


def m_expression(table: str, df: pd.DataFrame, source: str) -> list[str]:
    types = ", ".join(f'{{"{c}", {_dtype(df[c])[1]}}}' for c in df.columns)
    # A static base URL keeps the query refreshable in the Power BI Service
    content = (f'Web.Contents("https://raw.githubusercontent.com", [RelativePath = DataPath & "{table}.csv"])'
               if source == "github" else f'File.Contents(DataPath & "{table}.csv")')
    return [
        "let",
        f'    Source = Csv.Document({content}, [Delimiter=",", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),',
        "    Promoted = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),",
        '    Blanks = Table.TransformColumns(Promoted, {}, each if _ = "" then null else _),',
        f'    Typed = Table.TransformColumnTypes(Blanks, {{{types}}}, "en-US")',
        "in",
        "    Typed",
    ]


MEASURES = [
    ("Users", "COUNTROWS(fact_user)", "#,0", "Experiment"),
    ("User-Days", "SUM(fact_daily[users_assigned])", "#,0", "Engagement"),
    ("Minutes per User-Day", "DIVIDE(SUM(fact_daily[minutes]), [User-Days])", "0.00", "Engagement"),
    ("DAU Rate", "DIVIDE(SUM(fact_daily[active_users]), [User-Days])", "0.0%", "Engagement"),
    ("Sessions per User-Day", "DIVIDE(SUM(fact_daily[sessions]), [User-Days])", "0.00", "Engagement"),
    ("Friend Interactions per User-Day", "DIVIDE(SUM(fact_daily[friend_interactions]), [User-Days])", "0.000", "Engagement"),
    ("Posts per User-Day", "DIVIDE(SUM(fact_daily[posts_created]), [User-Days])", "0.000", "Engagement"),
    ("Revenue per User-Day", "DIVIDE(SUM(fact_daily[ad_revenue]), [User-Days])", "$0.0000", "Monetisation"),
    ("Ad CTR", "DIVIDE(SUM(fact_daily[ad_clicks]), SUM(fact_daily[ad_impressions]))", "0.00%", "Monetisation"),
    ("Negative Feedback per 100 Min", "DIVIDE(SUM(fact_daily[negative_feedback]), SUM(fact_daily[minutes])) * 100", "0.000", "Guardrails"),
    ("Minutes per User-Day (Control)", 'CALCULATE([Minutes per User-Day], REMOVEFILTERS(dim_arm), dim_arm[arm] = "control")', "0.00", "Lift vs Control"),
    ("Minutes Lift vs Control", "DIVIDE([Minutes per User-Day], [Minutes per User-Day (Control)]) - 1", "+0.0%;-0.0%;0.0%", "Lift vs Control"),
    ("Revenue Lift vs Control", 'DIVIDE([Revenue per User-Day], CALCULATE([Revenue per User-Day], REMOVEFILTERS(dim_arm), dim_arm[arm] = "control")) - 1', "+0.0%;-0.0%;0.0%", "Lift vs Control"),
    ("Friend Interactions Lift vs Control", 'DIVIDE([Friend Interactions per User-Day], CALCULATE([Friend Interactions per User-Day], REMOVEFILTERS(dim_arm), dim_arm[arm] = "control")) - 1', "+0.0%;-0.0%;0.0%", "Lift vs Control"),
    ("Churn Rate", "AVERAGE(fact_user[churned])", "0.00%", "Retention"),
    ("Churn Rate (Control)", 'CALCULATE([Churn Rate], REMOVEFILTERS(dim_arm), dim_arm[arm] = "control")', "0.00%", "Retention"),
    ("Churn Change vs Control (pp)", "([Churn Rate] - [Churn Rate (Control)]) * 100", "+0.00;-0.00;0.00", "Retention"),
    ("Avg Churn Risk", "AVERAGE(fact_user[churn_probability])", "0.0%", "Retention"),
    ("High-Risk Users", 'CALCULATE(COUNTROWS(fact_user), fact_user[risk_tier] IN {"High", "Critical"})', "#,0", "Retention"),
    ("AI Feed Minutes Lift", 'CALCULATE(MAX(metric_results[rel_lift]), metric_results[arm] = "ai_feed", metric_results[metric] = "minutes")', "+0.0%;-0.0%", "Headline KPIs"),
    ("Friend Boost Minutes Lift", 'CALCULATE(MAX(metric_results[rel_lift]), metric_results[arm] = "ai_feed_social", metric_results[metric] = "minutes")', "+0.0%;-0.0%", "Headline KPIs"),
    ("AI Feed Churn Change", 'CALCULATE(MAX(metric_results[rel_lift]), metric_results[arm] = "ai_feed", metric_results[metric] = "churned")', "+0.0%;-0.0%", "Headline KPIs"),
    ("Friend Boost Churn Change", 'CALCULATE(MAX(metric_results[rel_lift]), metric_results[arm] = "ai_feed_social", metric_results[metric] = "churned")', "+0.0%;-0.0%", "Headline KPIs"),
    ("Friend Boost 12M Value ($M)", 'CALCULATE(MAX(impact_monte_carlo[p50_musd]), impact_monte_carlo[comparison] = "AI + friend boost vs control")', "$#,0\"M\"", "Headline KPIs"),
    ("AI Feed 12M Value ($M)", 'CALCULATE(MAX(impact_monte_carlo[p50_musd]), impact_monte_carlo[comparison] = "AI feed vs control")', "$#,0\"M\"", "Headline KPIs"),
    ("Share of Draws FB Beats AI", 'CALCULATE(MAX(impact_monte_carlo[prob_positive]), impact_monte_carlo[comparison] = "Friend boost minus AI feed")', "0.0%", "Headline KPIs"),
    ("Significant Wins", 'CALCULATE(COUNTROWS(metric_results), metric_results[verdict] = "win")', "0", "Headline KPIs"),
    ("Significant Harms", 'CALCULATE(COUNTROWS(metric_results), metric_results[verdict] = "harm")', "0", "Headline KPIs"),
    ("Best Policy Uplift ($/user/yr)", "MAX(policy_values[uplift_vs_control_usd])", "$0.00", "Headline KPIs"),
    ("Churn Model AUC", 'CALCULATE(MAX(churn_model_metrics[roc_auc]), churn_model_metrics[deployed] = 1)', "0.000", "Headline KPIs"),
]

RELATIONSHIPS = [
    ("fact_daily", "date", "dim_date", "date"), ("fact_daily", "arm", "dim_arm", "arm"),
    ("fact_user", "arm", "dim_arm", "arm"), ("metric_results", "arm", "dim_arm", "arm"),
    ("segment_results", "arm", "dim_arm", "arm"), ("daily_lift", "arm", "dim_arm", "arm"),
    ("km_survival", "arm", "dim_arm", "arm"), ("impact_fan", "arm", "dim_arm", "arm"),
    ("topic_exposure_by_arm", "arm", "dim_arm", "arm"), ("content_exposure_summary", "arm", "dim_arm", "arm"),
    ("sequential_monitoring", "arm", "dim_arm", "arm"), ("impact_distribution", "arm", "dim_arm", "arm"),
    ("qini_curves", "arm", "dim_arm", "arm"), ("cate_deciles", "arm", "dim_arm", "arm"),
    ("decision_guardrails", "arm", "dim_arm", "arm"),
]


def build_model(tables: dict[str, pd.DataFrame], data_path: str, source: str) -> dict:
    tmsl_tables = []
    for t, df in tables.items():
        cols = []
        for c in df.columns:
            dt, _ = _dtype(df[c])
            col = {"name": c, "dataType": dt, "sourceColumn": c, "summarizeBy": "none" if dt in ("string", "dateTime") or c.endswith("_id") else "sum",
                   "lineageTag": str(uuid.uuid5(uuid.NAMESPACE_URL, f"{t}.{c}"))}
            if c in DISPLAY_FORMATS:
                col["formatString"] = DISPLAY_FORMATS[c]
            elif dt == "dateTime":
                col["formatString"] = "d mmm"
            elif c in PCT_COLS or c.endswith("_lift") or c.endswith("_rel_change"):
                col["formatString"] = "0.0%"
            elif dt == "double":
                col["formatString"] = "#,0.0000"
            if t == "dim_arm" and c == "arm_label":
                col["sortByColumn"] = "sort_order"
            cols.append(col)
        tmsl_tables.append({"name": t, "lineageTag": str(uuid.uuid5(uuid.NAMESPACE_URL, t)), "columns": cols,
                            "partitions": [{"name": t, "mode": "import", "source": {"type": "m", "expression": m_expression(t, df, source)}}]})
    tmsl_tables.append({
        "name": "_Measures", "lineageTag": str(uuid.uuid5(uuid.NAMESPACE_URL, "_Measures")),
        "columns": [{"name": "Measure Holder", "dataType": "string", "sourceColumn": "Measure Holder", "isHidden": True,
                     "summarizeBy": "none", "lineageTag": str(uuid.uuid5(uuid.NAMESPACE_URL, "_Measures.holder"))}],
        "measures": [{"name": n, "expression": e, "formatString": f, "displayFolder": folder,
                      "lineageTag": str(uuid.uuid5(uuid.NAMESPACE_URL, f"measure.{n}"))} for n, e, f, folder in MEASURES],
        "partitions": [{"name": "_Measures", "mode": "import", "source": {"type": "m", "expression": [
            "let", '    Source = #table(type table [#"Measure Holder" = text], {})', "in", "    Source"]}}],
    })
    rels = [{"name": str(uuid.uuid5(uuid.NAMESPACE_URL, f"{a}.{b}->{c}.{d}")), "fromTable": a, "fromColumn": b,
             "toTable": c, "toColumn": d} for a, b, c, d in RELATIONSHIPS]
    return {
        "name": NAME, "compatibilityLevel": 1567,
        "model": {
            "culture": "en-US", "defaultPowerBIDataSourceVersion": "powerBI_V3", "sourceQueryCulture": "en-US",
            "dataAccessOptions": {"legacyRedirects": True, "returnErrorValuesAsNull": True},
            "tables": tmsl_tables, "relationships": rels,
            "expressions": [{"name": "DataPath", "kind": "m", "lineageTag": str(uuid.uuid5(uuid.NAMESPACE_URL, "DataPath")),
                             "expression": f'"{data_path}" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=true]'}],
            "annotations": [{"name": "PBI_QueryOrder", "value": json.dumps(["DataPath"] + list(tables) + ["_Measures"])},
                            {"name": "__PBI_TimeIntelligenceEnabled", "value": "0"}],
        },
    }




def theme() -> dict:
    return {"name": "Feed Ranking Launch Review", "dataColors": [C.ARM_COLORS["control"], C.ARM_COLORS["ai_feed"], C.ARM_COLORS["ai_feed_social"],
                                                    "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"],
            "background": "#FCFCFB", "foreground": "#0B0B0B", "tableAccent": "#2a78d6",
            "good": C.STATUS["good"], "neutral": "#898781", "bad": C.STATUS["critical"],
            "maximum": "#0d366b", "center": "#f0efec", "minimum": "#cde2fb",
            "textClasses": {"title": {"fontFace": "Segoe UI Semibold", "fontSize": 12, "color": "#0B0B0B"},
                            "label": {"fontFace": "Segoe UI", "fontSize": 9, "color": "#52514E"},
                            "callout": {"fontFace": "Segoe UI Semibold", "fontSize": 26, "color": "#0B0B0B"}},
            "visualStyles": {"*": {"*": {"background": [{"color": {"solid": {"color": "#FFFFFF"}}, "transparency": 0}],
                                         "border": [{"show": True, "color": {"solid": {"color": "#E1E0D9"}}, "radius": 6}]}}}}


def platform(item_type: str) -> str:
    """Fabric Git metadata. Without it a synced workspace ignores the folder."""
    return json.dumps({
        "$schema": "https://developer.microsoft.com/json-schemas/fabric/gitIntegration/platformProperties/2.0.0/schema.json",
        "metadata": {"type": item_type, "displayName": DISPLAY},
        "config": {"version": "2.0", "logicalId": str(uuid.uuid5(uuid.NAMESPACE_URL, f"{NAME}.{item_type}"))},
    }, indent=2)


def write_notebook():
    """Fabric notebook that creates or updates the model and report from this repo."""
    def md(t): return {"cell_type": "markdown", "metadata": {}, "source": t}
    def code(t): return {"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [], "source": t}
    repo = "https://raw.githubusercontent.com/" + C.GITHUB_DATA_PATH.rsplit("/data/", 1)[0]
    cells = [
        md(f"# {DISPLAY}\n\nCreates or updates the {DISPLAY} semantic model and report in this workspace from the GitHub "
           "repo. Run all (2-3 minutes); re-run whenever the repo changes."),
        code("%pip install -q semantic-link-labs"),
        code(f'''import time
import requests
import sempy.fabric as fabric
import sempy_labs as labs
from sempy_labs.report import create_report_from_reportjson, update_report_from_reportjson

REPO = "{repo}"
FOLDER = "{NAME}"
NAME = "{DISPLAY}"
nocache = {{"Cache-Control": "no-cache"}}

bim = requests.get(f"{{REPO}}/{{FOLDER}}.SemanticModel/model.bim", headers=nocache, timeout=60).json()
report_json = requests.get(f"{{REPO}}/{{FOLDER}}.Report/report.json", headers=nocache, timeout=60).json()
print(f"Loaded model: {{len(bim['model']['tables'])}} tables; report: {{len(report_json['sections'])}} pages")'''),
        md("## 1. Semantic model (create or update)"),
        code('''if NAME in fabric.list_datasets(mode="rest")["Dataset Name"].values:
    labs.update_semantic_model_from_bim(dataset=NAME, bim_file=bim)
else:
    labs.create_semantic_model_from_bim(dataset=NAME, bim_file=bim)
for _ in range(30):
    if NAME in fabric.list_datasets(mode="rest")["Dataset Name"].values:
        print("semantic model is ready")
        break
    time.sleep(5)'''),
        md("## 2. Load the data\n\nOn the first run this may fail on credentials. Set the Web source to Anonymous / Public "
           "under the semantic model's Settings → Data source credentials, then re-run this cell."),
        code("labs.refresh_semantic_model(dataset=NAME)"),
        md("## 3. Report (create or update)"),
        code('''if NAME in fabric.list_reports()["Name"].values:
    update_report_from_reportjson(report=NAME, report_json=report_json)
else:
    create_report_from_reportjson(report=NAME, dataset=NAME, report_json=report_json)
print(f"Done. Open the '{NAME}' report (refresh the browser tab if it was already open).")'''),
    ]
    nb = {"cells": cells, "metadata": {"kernelspec": {"name": "synapse_pyspark", "display_name": "Synapse PySpark"},
                                       "language_info": {"name": "python"}}, "nbformat": 4, "nbformat_minor": 5}
    (OUT / f"{NAME}.ipynb").write_text(json.dumps(nb, indent=1))


def write_pbip(tables, data_path: str, source: str):
    for old in glob.glob(str(OUT / "SEI_*")) + glob.glob(str(OUT / "deploy_to_fabric.ipynb")):  # old names
        shutil.rmtree(old) if os.path.isdir(old) else os.remove(old)
    model_dir, report_dir = OUT / f"{NAME}.SemanticModel", OUT / f"{NAME}.Report"
    for d in (model_dir, report_dir):
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
    (OUT / f"{NAME}.pbip").write_text(json.dumps({"version": "1.0", "artifacts": [{"report": {"path": f"{NAME}.Report"}}],
                                                 "settings": {"enableAutoRecovery": True}}, indent=2))
    (model_dir / "definition.pbism").write_text(json.dumps({"version": "1.0", "settings": {}}, indent=2))
    (model_dir / "model.bim").write_text(json.dumps(build_model(tables, data_path, source), indent=2))
    (model_dir / ".platform").write_text(platform("SemanticModel"))
    (report_dir / ".platform").write_text(platform("Report"))
    (report_dir / "definition.pbir").write_text(json.dumps(
        {"version": "1.0", "datasetReference": {"byPath": {"path": f"../{NAME}.SemanticModel"}, "byConnection": None}}, indent=2))
    (report_dir / "report.json").write_text(json.dumps(build_report(facts()), indent=2))
    (OUT / f"{NAME}_Theme.json").write_text(json.dumps(theme(), indent=2))
    write_notebook()
    (OUT / ".gitignore").write_text("**/.pbi/localSettings.json\n**/.pbi/cache.abf\n")
    dax = [f"// DAX measures for the {DISPLAY} semantic model (table: _Measures)\n"]
    for n, e, f, folder in MEASURES:
        dax.append(f"// [{folder}]  format: {f}\n{n} =\n    {e}\n")
    (OUT / "measures.dax").write_text("\n".join(dax))


def write_excel(tables: dict[str, pd.DataFrame]):
    """All tables in one workbook, for uploading to the Power BI web app."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    path = OUT / f"{NAME}_data.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for t, df in tables.items():
            df.to_excel(xw, sheet_name=t, index=False)
            ws = xw.sheets[t]
            ref = f"A1:{get_column_letter(df.shape[1])}{len(df) + 1}"
            tab = Table(displayName=t, ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(tab)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", choices=["github", "local"], default="github",
                    help="github: read CSVs from the public repo (works in Desktop AND the Power BI Service); "
                         "local: read from a folder on disk (Desktop only)")
    ap.add_argument("--github-path", default=C.GITHUB_DATA_PATH, help="owner/repo/branch/path/ to powerbi/data/")
    ap.add_argument("--data-folder", default="C:\\SEI\\powerbi\\data\\",
                    help="absolute path (with trailing slash) to powerbi/data, used with --source local")
    args = ap.parse_args()
    if args.source == "github":
        data_path = args.github_path if args.github_path.endswith("/") else args.github_path + "/"
    else:
        data_path = args.data_folder if args.data_folder.endswith(("\\", "/")) else args.data_folder + "\\"
    tables = export_data()
    write_pbip(tables, data_path, args.source)
    xlsx = write_excel(tables)
    print(f"wrote {xlsx.name} for the Power BI web app")
    total = sum(len(t) for t in tables.values())
    print(f"exported {len(tables)} tables ({total:,} rows) to {DATA} and wrote {NAME}.pbip "
          f"(source = {args.source}, DataPath = {data_path})")


if __name__ == "__main__":
    main()
